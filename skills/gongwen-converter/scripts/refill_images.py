#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把源 docx 中的内嵌图片，按位置回填到 Pandoc 转换后的公文 docx 中。

问题背景
--------
`convert.py` 走的是 Markdown → docx，Markdown 里没有图片时，
源文档中的内嵌图片（思维导图、插图、印章等）会被整体丢弃。
本脚本以「图片所在段落之前最近的非空文本段落」作为锚点，
在目标 docx 中定位同一锚点段落，并把图片插入其后。

用法
----
    python refill_images.py <源.docx> <目标.docx> [-w 15.6]

    -w  图片宽度（cm），默认 15.6（A4 公文版心宽度 = 21 - 2.8 - 2.6）
        按原始宽高比自动推算高度。

说明
----
- 仅处理正文（word/document.xml）中的图片，页眉/页脚图片不处理。
- 锚点匹配忽略空格；找不到对应锚点时会打印警告并跳过该图，
  不会破坏目标文档。
- 插入的图片段落会显式设 `lineRule="auto"`，覆盖公文样式的
  「固定行距 28 磅」，否则图片会被裁切。
"""

import argparse
import os
import re
import shutil
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"

MEDIA_RE = re.compile(r"^word/media/")


def _norm(s):
    return re.sub(r"\s+", "", s or "")


def extract_source_images(src_docx, workdir):
    """返回 [(锚点文本, 图片绝对路径, (cx, cy) or None), ...]，按正文出现顺序。"""
    with zipfile.ZipFile(src_docx) as z:
        names = z.namelist()
        media = [n for n in names if MEDIA_RE.match(n) and not n.endswith("/")]
        for m in media:
            dest = os.path.join(workdir, os.path.basename(m))
            with open(dest, "wb") as f:
                f.write(z.read(m))

        rels = {}
        if "word/_rels/document.xml.rels" in names:
            rel_xml = z.read("word/_rels/document.xml.rels").decode("utf-8")
            for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rel_xml):
                rid, target = m.group(1), m.group(2)
                if target.startswith("media/"):
                    rels[rid] = os.path.join(workdir, os.path.basename(target))
        doc_xml = z.read("word/document.xml").decode("utf-8")

    paras = re.findall(r"<w:p[ >].*?</w:p>|<w:p/>", doc_xml, re.S)
    result = []
    anchor = ""
    for p in paras:
        txt = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", p, re.S))
        if _norm(txt):
            anchor = txt
        if "<w:drawing>" not in p and "<w:pict>" not in p:
            continue
        exts = re.findall(r'<wp:extent cx="(\d+)" cy="(\d+)"', p)
        extent = (int(exts[0][0]), int(exts[0][1])) if exts else None
        for rid in re.findall(r'r:embed="([^"]+)"', p) + re.findall(r'r:id="([^"]+)"', p):
            path = rels.get(rid)
            if path and os.path.exists(path):
                result.append((anchor, path, extent))
    return result


def refill(src_docx, dst_docx, width_cm=15.6):
    workdir = tempfile.mkdtemp(prefix="refill_img_")
    try:
        images = extract_source_images(src_docx, workdir)
        if not images:
            print("[INFO] 源文档正文中没有图片")
            return 0

        doc = Document(dst_docx)
        inserted = 0
        for anchor_txt, img_path, _extent in images:
            key = _norm(anchor_txt)
            anchor_para = None
            for p in doc.paragraphs:
                if _norm(p.text) == key:
                    anchor_para = p
                    break
            if anchor_para is None:
                print(f"[WARN] 目标文档未找到锚点段落，跳过该图片: {anchor_txt[:24]!r}")
                continue

            pic_p = doc.add_paragraph()
            pf = pic_p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.first_line_indent = 0
            pf.line_spacing = 1.0  # 覆盖样式里的「固定行距」，否则图片被裁切
            pf.space_before = 0
            pf.space_after = 0
            pic_p.add_run().add_picture(img_path, width=Cm(width_cm))
            anchor_para._element.addnext(pic_p._element)
            inserted += 1
            print(f"[OK] 已在 {anchor_txt[:24]!r} 之后插入图片: {os.path.basename(img_path)}")

        if inserted:
            doc.save(dst_docx)
        print(f"[完成] 共回填 {inserted} 张图片 → {dst_docx}")
        return inserted
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def main():
    ap = argparse.ArgumentParser(description="回填源 docx 中的图片到目标 docx")
    ap.add_argument("src", help="源 docx（含图片）")
    ap.add_argument("dst", help="目标 docx（转换结果）")
    ap.add_argument("-w", "--width", type=float, default=15.6, help="图片宽度 cm，默认 15.6")
    a = ap.parse_args()
    if not os.path.exists(a.src):
        print(f"[错误] 文件不存在: {a.src}")
        sys.exit(1)
    if not os.path.exists(a.dst):
        print(f"[错误] 文件不存在: {a.dst}")
        sys.exit(1)
    n = refill(a.src, a.dst, a.width)
    if n == 0:
        sys.exit(0)


if __name__ == "__main__":
    main()
