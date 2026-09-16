#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验生成的公文 docx 是否合规（GB/T 9704-2012）。

用法:
    python check_gongwen.py <目标.docx> [-o 报告.txt]

输出内容:
  1. 页面设置（A4 21×29.7cm，上 3.7 / 下 3.5 / 左 2.8 / 右 2.6 cm）
  2. 逐段样式清单（Title / Heading 1-4 / First Paragraph / Body Text）
  3. 样式级的字体、字号、固定行距、首行缩进，并检查是否残留主题字体引用
     （theme 引用会覆盖公文要求的黑体/楷体/仿宋，是常见故障点）
  4. 表格校验：全框线、表宽=版心宽、居中、表头跨页重复、垂直居中、
     表头黑体/表内仿宋五号、表内固定行距 16 磅
     （对应 fix_gongwen_format.py 的表格美化结果）

退出码: 0=全部通过, 2=存在问题
"""

import argparse
import sys

from docx import Document
from docx.oxml.ns import qn

# 表格目标值（与 fix_gongwen_format.py 保持一致）
TBL_W = 8844          # 版心宽 twip
TBL_EA_HEAD = "黑体"
TBL_EA_BODY = "仿宋_GB2312"
TBL_SZ = 21           # 五号 = 10.5pt = 21 半点
TBL_LINE = "320"      # 固定行距 16 磅 = 320 twip

CHECK_STYLES = [
    "Normal", "First Paragraph", "Body Text", "Title",
    "Heading 1", "Heading 2", "Heading 3", "Heading 4",
]

EXPECT_STYLE = {
    "Normal": ("仿宋_GB2312", 16.0),
    "First Paragraph": ("仿宋_GB2312", 16.0),
    "Body Text": ("仿宋_GB2312", 16.0),
    "Title": ("方正小标宋简体", 22.0),
    "Heading 1": ("黑体", 16.0),
    "Heading 2": ("楷体_GB2312", 16.0),
    "Heading 3": ("仿宋_GB2312", 16.0),
    "Heading 4": ("仿宋_GB2312", 16.0),
}


def run(docx_path):
    doc = Document(docx_path)
    out = []
    ok = True

    sec = doc.sections[0]
    out.append("== 页面设置 ==")
    out.append("  %.1f x %.1f cm | 上%.2f 下%.2f 左%.2f 右%.2f cm" % (
        sec.page_width.cm, sec.page_height.cm, sec.top_margin.cm,
        sec.bottom_margin.cm, sec.left_margin.cm, sec.right_margin.cm))
    out.append("")

    out.append("== 逐段结构 ==")
    out.append("  （flc=首行缩进字符数×100；段落未显式设置时标注 (样式) 表示继承自段落样式）")
    for i, p in enumerate(doc.paragraphs):
        st = p.style.name
        pPr = p._element.find(qn("w:pPr"))
        ind = pPr.find(qn("w:ind")) if pPr is not None else None
        flc = ind.get(qn("w:firstLineChars")) if ind is not None else None
        if flc is None:
            try:
                s_pPr = p.style.element.find(qn("w:pPr"))
                s_ind = s_pPr.find(qn("w:ind")) if s_pPr is not None else None
                s_flc = s_ind.get(qn("w:firstLineChars")) if s_ind is not None else None
            except Exception:
                s_flc = None
            flc = ("%s(样式)" % s_flc) if s_flc is not None else "-"
        draw = "图" if p._element.findall(".//" + qn("w:drawing")) else ""
        txt = p.text.strip()
        if not txt and not draw:
            out.append("[%02d] %-16s <空行>" % (i, st))
        else:
            out.append("[%02d] %-16s flc=%-10s %s%s" % (i, st, flc, draw, txt[:42]))
    out.append("")

    out.append("== 样式校验 ==")
    for name in CHECK_STYLES:
        try:
            s = doc.styles[name]
        except KeyError:
            out.append("  %-16s 缺失 ✗" % name)
            ok = False
            continue
        rPr = s.element.find(qn("w:rPr"))
        rfonts = rPr.find(qn("w:rFonts")) if rPr is not None else None
        ea = rfonts.get(qn("w:eastAsia")) if rfonts is not None else None
        asc = rfonts.get(qn("w:ascii")) if rfonts is not None else None
        theme = [a for a in ("asciiTheme", "eastAsiaTheme", "hAnsiTheme", "cstheme")
                 if rfonts is not None and rfonts.get(qn("w:" + a))]
        sz_el = rPr.find(qn("w:sz")) if rPr is not None else None
        sz = int(sz_el.get(qn("w:val"))) / 2.0 if sz_el is not None else None
        pPr = s.element.find(qn("w:pPr"))
        sp = pPr.find(qn("w:spacing")) if pPr is not None else None
        line = sp.get(qn("w:line")) if sp is not None else None
        rule = sp.get(qn("w:lineRule")) if sp is not None else None
        ind = pPr.find(qn("w:ind")) if pPr is not None else None
        flc = ind.get(qn("w:firstLineChars")) if ind is not None else None

        marks = []
        exp_ea, exp_sz = EXPECT_STYLE.get(name, (None, None))
        if exp_ea and ea != exp_ea:
            marks.append("字体应为 %s" % exp_ea)
        if exp_sz and sz != exp_sz:
            marks.append("字号应为 %s" % exp_sz)
        if line != "560" or rule != "exact":
            marks.append("行距应为固定 28 磅")
        if theme:
            marks.append("残留主题字体引用: %s" % theme)
        if name != "Title" and flc != "200":
            marks.append("首行缩进应为 2 字符")
        if marks:
            ok = False
        out.append("  %-16s ea=%-14s ascii=%-16s sz=%-5s flc=%-5s line=%-4s %-5s %s" % (
            name, ea, asc, sz, flc, line, rule, "✓" if not marks else "✗ " + "；".join(marks)))

    out.append("")
    out.append("== 表格校验 ==")
    if not doc.tables:
        out.append("  （文档无表格）")
    for ti, tbl in enumerate(doc.tables):
        marks = []
        el = tbl._element
        tblPr = el.find(qn("w:tblPr"))

        borders = tblPr.find(qn("w:tblBorders")) if tblPr is not None else None
        if borders is None:
            marks.append("缺全框线 tblBorders")
        else:
            missing = [e for e in ("top", "left", "bottom", "right",
                                   "insideH", "insideV")
                       if borders.find(qn("w:" + e)) is None]
            if missing:
                marks.append("框线不全: %s" % missing)

        tw = tblPr.find(qn("w:tblW")) if tblPr is not None else None
        got_w = tw.get(qn("w:w")) if tw is not None else None
        if got_w != str(TBL_W):
            marks.append("表宽应为 %d twip，实为 %s" % (TBL_W, got_w))

        jc = tblPr.find(qn("w:jc")) if tblPr is not None else None
        if jc is None or jc.get(qn("w:val")) != "center":
            marks.append("表格应居中")

        rows = tbl.rows
        head_tr = rows[0]._element.find(qn("w:trPr")) if rows else None
        if head_tr is None or head_tr.find(qn("w:tblHeader")) is None:
            marks.append("表头行缺 tblHeader（跨页不重复）")

        for ri, row in enumerate(rows):
            want_ea = TBL_EA_HEAD if ri == 0 else TBL_EA_BODY
            for ci, cell in enumerate(row.cells):
                tcPr = cell._element.find(qn("w:tcPr"))
                va = tcPr.find(qn("w:vAlign")) if tcPr is not None else None
                if va is None or va.get(qn("w:val")) != "center":
                    marks.append("第%d行第%d列未垂直居中" % (ri + 1, ci + 1))
                    break
                for p in cell.paragraphs:
                    if not p.text.strip():
                        continue
                    rPr = p.runs[0]._element.find(qn("w:rPr")) if p.runs else None
                    rf = rPr.find(qn("w:rFonts")) if rPr is not None else None
                    ea = rf.get(qn("w:eastAsia")) if rf is not None else None
                    sz_el = rPr.find(qn("w:sz")) if rPr is not None else None
                    sz = sz_el.get(qn("w:val")) if sz_el is not None else None
                    sp = p._element.find(qn("w:pPr"))
                    sp = sp.find(qn("w:spacing")) if sp is not None else None
                    ln = sp.get(qn("w:line")) if sp is not None else None
                    if ea != want_ea:
                        marks.append("第%d行字体应为 %s，实为 %s" % (ri + 1, want_ea, ea))
                    if sz != str(TBL_SZ):
                        marks.append("第%d行字号应为五号(%d)" % (ri + 1, TBL_SZ))
                    if ln != TBL_LINE:
                        marks.append("表内行距应为固定 16 磅")
                    break
                else:
                    continue
                break

        ncol = len(tbl.columns)
        grid = el.find(qn("w:tblGrid"))
        cols = [c.get(qn("w:w")) for c in grid.findall(qn("w:gridCol"))] if grid is not None else []
        out.append("  表%d: %d行 × %d列 | 列宽 %s | %s"
                   % (ti + 1, len(rows), ncol, cols,
                      "✓" if not marks else "✗ " + "；".join(marks)))
        if marks:
            ok = False

    out.append("")
    out.append("结论: " + ("全部通过 ✓" if ok else "存在问题，见上文 ✗"))
    return "\n".join(out), ok


def main():
    ap = argparse.ArgumentParser(description="校验公文 docx 格式")
    ap.add_argument("docx")
    ap.add_argument("-o", "--out", help="报告输出到 txt")
    a = ap.parse_args()
    report, ok = run(a.docx)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(report)
        print("报告已写入:", a.out)
    print(report)
    sys.exit(0 if ok else 2)


if __name__ == "__main__":
    main()
