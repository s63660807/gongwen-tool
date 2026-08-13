#!/usr/bin/env python3
"""
doc / docx 文件内容读取工具
用于写材料时提取已有文档内容。

用法:
    python read_doc.py <文件路径>            # 打印到屏幕
    python read_doc.py <文件路径> -o 输出.txt  # 保存到 txt 文件

支持:
    .docx  — 使用 python-docx 读取段落 + 表格
    .doc   — 优先使用 antiword（Git Bash 自带），失败则尝试 Word COM 转换
"""

import sys
import os
import subprocess
import shutil

HEADING_STYLES = {'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'}


def read_docx(path):
    """读取 .docx：输出段落（含标题层级）与表格"""
    from docx import Document
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = para.style.name if para.style else 'Normal'
        if style == 'Title':
            lines.append(f"【标题】{text}")
        elif style in ('Heading 1', 'Heading 2', 'Heading 3', 'Heading 4'):
            level = style.split()[-1]
            lines.append(f"{'#' * int(level)} {text}")
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def find_antiword():
    """查找 antiword：PATH > Git Bash 常见位置"""
    found = shutil.which("antiword")
    if found:
        return found
    # Git Bash 自带 antiword 的常见位置
    common_paths = [
        os.path.join(os.environ.get("USERPROFILE", ""), ".workbuddy-ai", "vendor", "PortableGit", "mingw64", "bin", "antiword.exe"),
        os.path.join(os.environ.get("USERPROFILE", ""), "scoop", "apps", "git", "current", "mingw64", "bin", "antiword.exe"),
        r"C:\Program Files\Git\mingw64\bin\antiword.exe",
    ]
    for p in common_paths:
        if os.path.exists(p):
            return p
    return None


def read_doc_antiword(path):
    """读取 .doc：使用 antiword（老格式 Word 文档）
    antiword 不支持中文路径，先复制到临时 ASCII 路径再读取"""
    antiword = find_antiword()
    if not antiword:
        return None
    import tempfile
    try:
        # 复制到纯 ASCII 临时路径（antiword 无法处理中文路径）
        tmpdir = tempfile.mkdtemp(prefix="docread_")
        tmp_path = os.path.join(tmpdir, "input.doc")
        shutil.copy2(path, tmp_path)
        result = subprocess.run(
            [antiword, "-w", "0", tmp_path],
            capture_output=True, text=True, timeout=30,
            encoding="utf-8", errors="replace"
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout.strip()
    except (subprocess.TimeoutExpired, OSError):
        pass
    finally:
        try:
            shutil.rmtree(tmpdir)
        except Exception:
            pass
    return None


def read_doc_com(path):
    """读取 .doc：使用 Word COM（Windows + 已安装 Microsoft Word）"""
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(path, ReadOnly=True)
        text = doc.Content.Text
        doc.Close(False)
        word.Quit()
        return text
    except Exception:
        try:
            word.Quit()
        except Exception:
            pass
        return None


def read_doc(path):
    """读取 .doc：antiword 优先，失败回退 Word COM"""
    text = read_doc_antiword(path)
    if text:
        return text
    text = read_doc_com(path)
    if text:
        return text
    raise RuntimeError(
        "无法读取 .doc 文件：antiword 解析失败，且未找到可用的 Microsoft Word"
    )


def main():
    if len(sys.argv) < 2:
        print("doc/docx 读取工具")
        print()
        print("用法:")
        print("    python read_doc.py <文件路径>               # 打印到屏幕")
        print("    python read_doc.py <文件路径> -o 输出.txt    # 保存到文件")
        sys.exit(0)

    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"[错误] 文件不存在: {path}")
        sys.exit(1)

    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        text = read_docx(path)
    elif ext == ".doc":
        text = read_doc(path)
    else:
        print(f"[错误] 不支持的文件类型: {ext}（仅支持 .doc / .docx）")
        sys.exit(1)

    # 输出到文件或屏幕
    if "-o" in sys.argv:
        idx = sys.argv.index("-o")
        out_path = sys.argv[idx + 1] if len(sys.argv) > idx + 1 else None
        if not out_path:
            print("[错误] -o 后需要指定输出文件路径")
            sys.exit(1)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"[完成] 已保存: {out_path}")
    else:
        print(text)


if __name__ == "__main__":
    main()
